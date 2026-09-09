#!/usr/bin/env python3
"""Read-only DOCX postcheck for report structure and forbidden dynamic content."""
from __future__ import annotations
import argparse
import re
from pathlib import Path
from docx import Document

FORBIDDEN = (
    "证据截图", "证据：artifacts/", "测试团队：", "报告生成日期：", "执行摘要", "渗透路径", "阶段总结", "综述",
    # 2026-09-09 起：报告只面向漏洞审核员，禁止已删除章节与项目/本地内部术语
    "限制与观察", "环境准备", "TARGET_APPID=", "TARGET_VERSION=",
    "auth_sessions", "本地凭证文件", "engagements/", "runs/", "操作员",
    "【请补充资产归属证明网址】", "【请补充备案系统证明网址】",
)
LONG_TITLE = re.compile(r"攻防成果报告.*（.*(?:https?://|\\；|；).*[）)]")
MAX_DESCRIPTION_CHARS = 200  # 成果描述硬上限；策划目标为 ≤100 字、2~3 句

def doc_text(path: Path) -> str:
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    parts.extend(c.text for table in doc.tables for row in table.rows for c in row.cells)
    return "\n".join(parts)

def check(path: Path) -> list[str]:
    doc = Document(path)
    paragraph_text = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    text = "\n".join(paragraph_text + [c.text for table in doc.tables for row in table.rows for c in row.cells])
    violations = [f"forbidden text: {item}" for item in FORBIDDEN if item in text]
    if LONG_TITLE.search(text): violations.append("dynamic long title detected")
    if "攻防成果报告" not in text: violations.append("template title missing")
    if not ("复现命令" in text or "验证步骤" in text): violations.append("reproduction section missing")
    if "一、目标信息" in text: violations.append("target information section must be omitted")
    if "序号" in text: violations.append("serial-number row must be omitted")
    if "成果1：" in text: violations.append("single-result prefix must be omitted")
    if "权限/角色" in text: violations.append("permission/role row must be omitted")
    if "涉及数据量" in text: violations.append("data-volume label must be replaced by impact scope")
    for section in ("三、存在问题", "四、整改建议"):
        if section in paragraph_text:
            start = paragraph_text.index(section) + 1
            end = paragraph_text.index("四、整改建议") if section == "三、存在问题" and "四、整改建议" in paragraph_text else len(paragraph_text)
            lines = paragraph_text[start:end]
            if len(lines) > 2: violations.append(f"{section} has more than two entries")
            if any(len(line) > 50 for line in lines): violations.append(f"{section} entry exceeds 50 characters")
    for table in doc.tables:
        for row in table.rows:
            cells = row.cells
            if len(cells) >= 2 and cells[0].text.strip() == "成果描述" and len(cells[1].text.strip()) > MAX_DESCRIPTION_CHARS:
                violations.append(f"成果描述 exceeds {MAX_DESCRIPTION_CHARS} characters; keep it concise (target ≤100 chars)")
    return violations

def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("docx", type=Path)
    args = parser.parse_args()
    violations = check(args.docx)
    if violations:
        for item in violations: print(f"[!] {item}")
        return 1
    print(f"[+] DOCX postcheck passed: {args.docx}")
    return 0

if __name__ == "__main__":
    raise SystemExit(main())
