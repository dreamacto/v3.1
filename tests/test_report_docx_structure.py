"""2026-09-09 报告硬规则：面向漏洞审核员，无内部章节/术语，命令可复现。"""
import json
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from docx import Document

from report_docx import build_report
from scripts.maintenance import postcheck


def base_finding() -> list[dict]:
    return [{
        "system": "测试系统",
        "url": "https://example.test/api/user",
        "vulnerability_family": "未授权访问",
        "description": "匿名请求即可读取敏感字段，未做任何鉴权校验",
        "steps": ["步骤一：直接 GET 该接口"],
        "commands": [{"cmd": "curl -s https://example.test/api/user", "note": "返回 200 与敏感字段"}],
    }]


def render(meta: dict, tmp: str) -> Path:
    out = Path(tmp) / "report.docx"
    return build_report(meta, base_finding(), out, {}, None)


def doc_text(path: Path) -> str:
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    parts.extend(c.text for t in doc.tables for r in t.rows for c in r.cells)
    return "\n".join(parts)


class ReportDocxStructureTests(unittest.TestCase):
    def test_no_limitations_section_even_when_meta_carries_limitations(self):
        with tempfile.TemporaryDirectory() as tmp:
            text = doc_text(render({"target_name": "测试目标", "limitations": ["内部限制说明A"]}, tmp))
            self.assertNotIn("限制与观察", text)
            self.assertNotIn("内部限制说明A", text)

    def test_no_env_prepare_section_even_when_meta_carries_env_lines(self):
        with tempfile.TemporaryDirectory() as tmp:
            text = doc_text(render({"target_name": "测试目标", "env_lines": ["TARGET_APPID=com.example.test"]}, tmp))
            self.assertNotIn("环境准备", text)
            self.assertNotIn("TARGET_APPID=", text)

    def test_meta_proofs_render_empty_without_placeholder(self):
        with tempfile.TemporaryDirectory() as tmp:
            text = doc_text(render({"target_name": "测试目标"}, tmp))
            self.assertIn("资产归属证明网址", text)
            self.assertIn("备案系统证明网址", text)
            self.assertNotIn("【请补充资产归属证明网址】", text)
            self.assertNotIn("【请补充备案系统证明网址】", text)

    def test_reproduction_commands_rendered(self):
        with tempfile.TemporaryDirectory() as tmp:
            text = doc_text(render({"target_name": "测试目标"}, tmp))
            self.assertIn("curl -s https://example.test/api/user", text)
            self.assertIn("返回 200 与敏感字段", text)


class ReportPostcheckRulesTests(unittest.TestCase):
    def test_v11_style_report_passes(self):
        doc = Document()
        doc.add_paragraph("攻防成果报告", style="Title")
        doc.add_paragraph("资产归属证明网址")
        doc.add_paragraph("备案系统证明网址")
        doc.add_paragraph("一、成果说明", style="Heading 1")
        doc.add_paragraph("详细复现命令或操作步骤")
        doc.add_paragraph("# 复现环境: Git Bash + curl")
        doc.add_paragraph("curl -s https://example.test/api")
        doc.add_paragraph("# 实测原始输出: 200")
        doc.add_paragraph("三、存在问题", style="Heading 1")
        doc.add_paragraph("服务端未校验鉴权")
        doc.add_paragraph("四、整改建议", style="Heading 1")
        doc.add_paragraph("补充接口鉴权")
        table = doc.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        table.rows[0].cells[0].text = "成果描述"
        table.rows[0].cells[1].text = "匿名可读取敏感字段"

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "ok.docx"
            doc.save(path)
            self.assertEqual(postcheck.check(path), [])

    def test_internal_terms_and_long_description_fail(self):
        doc = Document()
        doc.add_paragraph("攻防成果报告", style="Title")
        doc.add_paragraph("一、成果说明", style="Heading 1")
        doc.add_paragraph("五、限制与观察", style="Heading 1")
        doc.add_paragraph("环境准备")
        doc.add_paragraph("TARGET_APPID=com.example.test")
        doc.add_paragraph("证据档案：engagements/demo-20260901/artifacts/x.jsonl（本地凭证文件 auth_sessions.local.json）")
        table = doc.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        table.rows[0].cells[0].text = "成果描述"
        table.rows[0].cells[1].text = "长" * 201

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "bad.docx"
            doc.save(path)
            violations = "\n".join(postcheck.check(path))
            for expected in ("限制与观察", "环境准备", "TARGET_APPID=", "engagements/", "auth_sessions",
                             "本地凭证文件", "成果描述 exceeds"):
                self.assertIn(expected, violations)


if __name__ == "__main__":
    unittest.main()
